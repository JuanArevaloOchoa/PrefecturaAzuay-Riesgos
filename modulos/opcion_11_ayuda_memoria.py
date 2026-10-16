from datetime import datetime
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from modulos.comun import cargar_bd_arcgis, ruta_en_reportes

def set_cell_shading(cell, color_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def crear_reporte_aptos(lista_tablas, titulo_reporte, periodo_str):
    hoy = datetime.now()
    FechaCreacion = hoy.strftime("%d/%m/%Y %H:%M")
    nombre_archivo = hoy.strftime("%y%m%d%H%M - Ayuda Memoria %H%M") + ".docx"
    ruta_final = str(ruta_en_reportes(nombre_archivo, subcarpeta="AyudaMemoria"))
    ruta_plantilla = r"D:\OneDrive - GAD Provincial del Azuay\0.1 Recursos\DocEnBlanco.docx"

    try:
        doc = Document(ruta_plantilla)
    except Exception:
        doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Aptos'
    font.size = Pt(12)
    font.color.rgb = RGBColor(40, 40, 40)

    p_fecha = doc.add_paragraph()
    p_fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_fecha = p_fecha.add_run(f"Fecha de Actualización: {FechaCreacion}")
    run_fecha.font.size = Pt(10)

    p_sec = doc.add_paragraph()
    p_sec.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sec.paragraph_format.space_after = Pt(0)
    run_sec = p_sec.add_run("SECRETARÍA DE INFRAESTRUCTURAS")
    run_sec.bold = True

    p_coor = doc.add_paragraph()
    p_coor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_coor.paragraph_format.space_after = Pt(10)
    run_coor = p_coor.add_run("COORDINACIÓN DE RIESGOS")
    run_coor.bold = True
    run_coor.font.color.rgb = RGBColor(128, 128, 128)

    p_memoria = doc.add_paragraph()
    p_memoria.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_memoria.paragraph_format.space_after = Pt(0)
    run_memoria = p_memoria.add_run("AYUDA MEMORIA DE EVENTOS ATENDIDOS")
    run_memoria.font.size = Pt(14)
    run_memoria.bold = True

    p_periodo = doc.add_paragraph()
    p_periodo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_periodo.paragraph_format.space_after = Pt(10)
    run_periodo = p_periodo.add_run(f"PERÍODO: {periodo_str}")
    run_periodo.bold = True

    p2 = doc.add_paragraph(
        "Ayuda memoria correspondiente a los eventos atendidos por la Secretaría de Infraestructuras en la provincia del Azuay para el período correspondiente. "
        "A continuación, se detallan los recursos utilizados y las categorías de eventos."
    )
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for titulo_tabla, datos_matriz in lista_tablas:
        p_tabla = doc.add_paragraph()
        run_tabla = p_tabla.add_run(titulo_tabla)
        run_tabla.bold = True
        run_tabla.font.name = 'Aptos'

        filas = len(datos_matriz)
        columnas = len(datos_matriz[0])
        tabla = doc.add_table(rows=filas, cols=columnas)
        tabla.style = 'Table Grid'

        for i, fila in enumerate(datos_matriz):
            celdas = tabla.rows[i].cells
            for j, valor in enumerate(fila):
                texto = str(valor).upper() if i == 0 else str(valor)
                p = celdas[j].paragraphs[0]
                run = p.add_run(texto)
                run.font.name = 'Aptos'
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_shading(celdas[j], "D9D9D9")
        doc.add_paragraph()

    num_siguiente = len(lista_tablas) + 1
    p_coes = doc.add_paragraph()
    run_coes = p_coes.add_run(f"{num_siguiente}. COES ACTIVADOS")
    run_coes.bold = True
    run_coes.font.name = 'Aptos'
    doc.add_paragraph()

    num_siguiente += 1
    p_res = doc.add_paragraph()
    run_res = p_res.add_run(f"{num_siguiente}. RESOLUCIONES EMITIDAS POR LA SNGR")
    run_res.bold = True
    run_res.font.name = 'Aptos'
    doc.add_paragraph()

    doc.save(ruta_final)
    print(f"\n[ÉXITO] Reporte generado correctamente en:\n{ruta_final}")

def ejecutar():
    try:
        df = cargar_bd_arcgis()
        if df is None or df.empty:
            print("❌ No se pudo cargar la BD de ArcGIS (o está vacía).")
            return
        df = df.copy()

        if 'fecha_de_visita' in df.columns:
            if pd.api.types.is_numeric_dtype(df['fecha_de_visita']):
                df['fecha_de_visita'] = pd.to_datetime(df['fecha_de_visita'], unit='ms', errors='coerce')
            else:
                df['fecha_de_visita'] = pd.to_datetime(df['fecha_de_visita'], errors='coerce')
        df['anio'] = df['fecha_de_visita'].dt.year

        maquinaria_pesada = ['trailers', 'tractores', 'excavadoras', 'cargadoras',
                             'retro_excavadoras', 'mini_cargadoras', 'volquetes', 'cisternas']
        maquinaria_liviana = ['camionetas', 'drones']
        personal_operativo = ['cuadrilla', 'operadores_y_conductores']
        personal_tecnico = ['personal_t_cnico']

        cols_existentes = [col for col in (maquinaria_pesada + maquinaria_liviana + personal_operativo + personal_tecnico) if col in df.columns]
        df[cols_existentes] = df[cols_existentes].fillna(0).apply(pd.to_numeric, errors='coerce').fillna(0)

        print("--- FILTROS DEL REPORTE DE RIESGOS ---")
        print("1. 2025\n2. 2026\n3. Fecha personalizada")
        opcion_fecha = input("Seleccione el período (1/2/3): ")

        if opcion_fecha == "1":
            df_fecha = df[df['anio'] == 2025]
            periodo_str = "2025"
        elif opcion_fecha == "2":
            df_fecha = df[df['anio'] == 2026]
            periodo_str = "2026"
        elif opcion_fecha == "3":
            fecha_inicio_str = input("Ingrese la fecha de inicio (YYYY-MM-DD): ")
            fecha_fin_str = input("Ingrese la fecha de fin (YYYY-MM-DD): ")
            try:
                fecha_inicio = pd.to_datetime(fecha_inicio_str)
                fecha_fin = pd.to_datetime(fecha_fin_str + ' 23:59:59')
                df_fecha = df[(df['fecha_de_visita'] >= fecha_inicio) & (df['fecha_de_visita'] <= fecha_fin)]
                periodo_str = f"DEL {fecha_inicio.strftime('%d-%m-%Y')} AL {fecha_fin.strftime('%d-%m-%Y')}"
            except Exception:
                df_fecha = df[df['anio'] == 2025]
                periodo_str = "2025"
        else:
            df_fecha = df[df['anio'] == 2025]
            periodo_str = "2025"

        col_tipo_atencion = 'tipo_de_atenci_n' if 'tipo_de_atenci_n' in df.columns else 'tipo_de_atencion'
        tipos_disponibles = sorted(df[col_tipo_atencion].dropna().unique().tolist())
        print("\nSeleccione el tipo de atención:\n0. Todos")
        for i, tipo in enumerate(tipos_disponibles, 1):
            print(f"{i}. {tipo}")

        opcion_tipo = input("Ingrese el número de su opción: ")
        if opcion_tipo in ["0", ""]:
            tipo_filtro = "Todos"
            df_filtrado = df_fecha
        else:
            try:
                idx = int(opcion_tipo) - 1
                tipo_filtro = tipos_disponibles[idx]
                df_filtrado = df_fecha[df_fecha[col_tipo_atencion] == tipo_filtro]
            except Exception:
                tipo_filtro = "Todos"
                df_filtrado = df_fecha

        if df_filtrado.empty:
            print(f"\nNo se encontraron datos para el período '{periodo_str}' y tipo: {tipo_filtro}")
            return

        maq_pesada_validas = [col for col in maquinaria_pesada if col in df.columns]
        maq_liviana_validas = [col for col in maquinaria_liviana if col in df.columns]
        pers_operativo_validos = [col for col in personal_operativo if col in df.columns]
        pers_tecnico_validos = [col for col in personal_tecnico if col in df.columns]

        resumen_final = {
            'Concepto': [
                'Cantidad Total de Eventos', 'Total Maquinaria Pesada',
                'Total Maquinaria Liviana', 'Total Operadores/Cuadrilla', 'Total Personal Técnico'
            ],
            'Valor': [
                len(df_filtrado),
                int(df_filtrado[maq_pesada_validas].values.sum()) if maq_pesada_validas else 0,
                int(df_filtrado[maq_liviana_validas].values.sum()) if maq_liviana_validas else 0,
                int(df_filtrado[pers_operativo_validos].values.sum()) if pers_operativo_validos else 0,
                int(df_filtrado[pers_tecnico_validos].values.sum()) if pers_tecnico_validos else 0
            ]
        }
        df_tabla_totales = pd.DataFrame(resumen_final)
        matriz_totales = [df_tabla_totales.columns.tolist()] + df_tabla_totales.values.tolist()

        col_evento = 'tipo_de_evento' if 'tipo_de_evento' in df.columns else 'tipo_evento'
        resumen_eventos = df_filtrado[col_evento].value_counts().reset_index()
        resumen_eventos.columns = ['Tipo de Evento', 'Cantidad']
        matriz_eventos = [resumen_eventos.columns.tolist()] + resumen_eventos.values.tolist()

        col_canton = 'cant_n' if 'cant_n' in df.columns else 'canton'
        resumen_geo = df_filtrado.groupby([col_canton, 'parroquia']).size().reset_index(name='No. DE REGISTROS')
        resumen_geo.columns = ['CANTÓN', 'PARROQUIA', 'No. DE REGISTROS']
        matriz_geo = [resumen_geo.columns.tolist()] + resumen_geo.values.tolist()

        lista_de_tablas = [
            ("1. RESUMEN DE CANTIDADES TOTALES", matriz_totales),
            ("2. CANTIDAD POR TIPO DE EVENTO", matriz_eventos),
            ("3. REGISTROS POR UBICACIÓN GEOGRÁFICA", matriz_geo)
        ]

        if tipo_filtro == "Todos":
            resumen_atencion = df_filtrado[col_tipo_atencion].value_counts().reset_index()
            resumen_atencion.columns = ['Tipo de Atención', 'Cantidad']
            matriz_atencion = [resumen_atencion.columns.tolist()] + resumen_atencion.values.tolist()
            lista_de_tablas.append(("4. CANTIDAD POR TIPO DE ATENCIÓN", matriz_atencion))

        titulo_doc = f"PERIODO: {periodo_str} - ATENCIÓN: {tipo_filtro}"
        crear_reporte_aptos(lista_de_tablas, titulo_doc, periodo_str)
    except Exception as e:
        print(f"Ocurrió un error al procesar los datos de ArcGIS: {e}")